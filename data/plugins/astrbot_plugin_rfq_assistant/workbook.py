"""Excel catalog import and reference-selection export helpers."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .domain import CATALOG_HEADERS, DEFAULT_PRODUCTS, FIELD_DEFINITIONS, MEDIUM_LABELS

DISCLAIMER = "演示参考价，需商务核价，不得直接作为正式报价"
_THIN_BORDER = Border(
    left=Side(style="thin", color="D9E0EA"),
    right=Side(style="thin", color="D9E0EA"),
    top=Side(style="thin", color="D9E0EA"),
    bottom=Side(style="thin", color="D9E0EA"),
)


def _catalog_row(product: dict[str, Any]) -> list[Any]:
    """Convert one product dictionary to the fixed spreadsheet order.

    Args:
        product: Validated product dictionary.

    Returns:
        Values matching ``CATALOG_HEADERS``.
    """
    values = []
    for key in CATALOG_HEADERS:
        value = product.get(key)
        if key == "compatible_media" and isinstance(value, list):
            value = ",".join(value)
        if key == "active":
            value = 1 if value else 0
        values.append(value)
    return values


def build_catalog_template(path: Path) -> Path:
    """Create an editable catalog template with the synthetic demo products.

    Args:
        path: Output XLSX path.

    Returns:
        The output path.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品目录"
    sheet.append(CATALOG_HEADERS)
    for product in DEFAULT_PRODUCTS:
        sheet.append(_catalog_row(product))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5EFF")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(CATALOG_HEADERS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = max(
            14, len(header) + 2
        )

    help_sheet = workbook.create_sheet("填写说明")
    help_sheet.append(["字段", "说明"])
    descriptions = {
        "compatible_media": "逗号分隔，仅支持 clean_water、mild_corrosive",
        "reference_price_cny": "仅作参考展示，插件不会计算税费、运费、折扣或总价",
        "active": "1 表示启用，0 表示停用",
    }
    for header in CATALOG_HEADERS:
        help_sheet.append([header, descriptions.get(header, "必填固定列")])
    help_sheet.column_dimensions["A"].width = 26
    help_sheet.column_dimensions["B"].width = 72
    workbook.save(path)
    return path


def read_catalog(path: Path) -> list[dict[str, Any]]:
    """Validate and read a fixed-schema product catalog.

    Args:
        path: Uploaded XLSX file.

    Returns:
        Validated active product dictionaries.

    Raises:
        ValueError: If columns or product values are invalid.
    """
    workbook = load_workbook(path, data_only=True, read_only=True)
    if "产品目录" not in workbook.sheetnames:
        raise ValueError("Excel 必须包含“产品目录”工作表")
    sheet = workbook["产品目录"]
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [
            str(value).strip() if value is not None else "" for value in next(rows)
        ]
    except StopIteration as exc:
        raise ValueError("产品目录为空") from exc
    if headers != CATALOG_HEADERS:
        raise ValueError("产品目录列与模板不一致，请先下载最新模板")

    numeric_fields = {
        "flow_min_m3h",
        "flow_max_m3h",
        "head_min_m",
        "head_max_m",
        "temp_min_c",
        "temp_max_c",
        "motor_power_kw",
        "voltage_v",
        "frequency_hz",
        "reference_price_cny",
        "lead_time_days",
    }
    products = []
    seen_skus = set()
    for row_number, row in enumerate(rows, 2):
        if all(value in (None, "") for value in row):
            continue
        product = dict(zip(CATALOG_HEADERS, row, strict=False))
        sku = str(product.get("sku") or "").strip()
        if not sku:
            raise ValueError(f"第 {row_number} 行缺少 sku")
        if sku in seen_skus:
            raise ValueError(f"sku 重复：{sku}")
        seen_skus.add(sku)
        product["sku"] = sku
        product["model_name"] = str(product.get("model_name") or "").strip()
        if not product["model_name"]:
            raise ValueError(f"第 {row_number} 行缺少 model_name")
        for key in numeric_fields:
            try:
                product[key] = float(product[key])
            except (TypeError, ValueError) as exc:
                raise ValueError(f"第 {row_number} 行的 {key} 必须是数字") from exc
        if product["flow_min_m3h"] > product["flow_max_m3h"]:
            raise ValueError(f"第 {row_number} 行流量下限大于上限")
        if product["head_min_m"] > product["head_max_m"]:
            raise ValueError(f"第 {row_number} 行扬程下限大于上限")
        if product["temp_min_c"] > product["temp_max_c"]:
            raise ValueError(f"第 {row_number} 行温度下限大于上限")
        if product["reference_price_cny"] < 0 or product["lead_time_days"] < 0:
            raise ValueError(f"第 {row_number} 行价格和交期不能为负数")
        media = [
            value.strip()
            for value in str(product.get("compatible_media") or "").split(",")
            if value.strip()
        ]
        if not media or any(value not in MEDIUM_LABELS for value in media):
            raise ValueError(
                f"第 {row_number} 行 compatible_media 仅支持 clean_water、mild_corrosive"
            )
        product["compatible_media"] = media
        product["wetted_material"] = str(product.get("wetted_material") or "").strip()
        product["ip_rating"] = str(product.get("ip_rating") or "").strip()
        product["seal_type"] = str(product.get("seal_type") or "").strip()
        product["price_note"] = str(product.get("price_note") or "").strip()
        active_value = product.get("active")
        product["active"] = str(active_value).strip().lower() not in {
            "0",
            "false",
            "no",
            "否",
            "none",
            "",
        }
        if product["active"]:
            products.append(product)
    if not products:
        raise ValueError("产品目录中至少需要一个启用型号")
    return products


def _style_table(sheet, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    """Apply consistent borders and wrapping to a worksheet range.

    Args:
        sheet: OpenPyXL worksheet.
        min_row: First row index.
        max_row: Last row index.
        min_col: First column index.
        max_col: Last column index.
    """
    for row in sheet.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_case_export(path: Path, case: dict[str, Any]) -> Path:
    """Create the approved reference-selection workbook.

    Args:
        path: Output XLSX path.
        case: Full case dictionary from storage.

    Returns:
        The generated output path.

    Raises:
        ValueError: If the case has no recommendation snapshot.
    """
    recommendation = case.get("recommendation")
    if not recommendation or not recommendation.get("candidates"):
        raise ValueError("当前询价没有可导出的候选方案")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    requirement_sheet = workbook.active
    requirement_sheet.title = "需求确认"
    requirement_sheet.merge_cells("A1:E1")
    requirement_sheet["A1"] = DISCLAIMER
    requirement_sheet["A1"].font = Font(bold=True, color="B42318", size=13)
    requirement_sheet["A1"].fill = PatternFill("solid", fgColor="FEE4E2")
    requirement_sheet["A1"].alignment = Alignment(horizontal="center")
    requirement_sheet.append(["询价编号", case["id"], "当前状态", case["status"], ""])
    requirement_sheet.append(["字段", "确认值", "单位", "来源", "原文证据"])
    for key, definition in FIELD_DEFINITIONS.items():
        field = case["revision"]["requirements"].get(key, {})
        value = field.get("value")
        if key == "medium":
            value = MEDIUM_LABELS.get(value, value)
        if isinstance(value, list):
            value = "；".join(str(item) for item in value)
        requirement_sheet.append(
            [
                definition["label"],
                value if value is not None else "未提供",
                field.get("unit", ""),
                field.get("origin", ""),
                field.get("evidence", ""),
            ]
        )
    for cell in requirement_sheet[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="344054")
    _style_table(requirement_sheet, 2, requirement_sheet.max_row, 1, 5)
    requirement_sheet.column_dimensions["A"].width = 20
    requirement_sheet.column_dimensions["B"].width = 28
    requirement_sheet.column_dimensions["C"].width = 12
    requirement_sheet.column_dimensions["D"].width = 14
    requirement_sheet.column_dimensions["E"].width = 52

    comparison_sheet = workbook.create_sheet("方案对比")
    candidates = recommendation["candidates"]
    comparison_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    comparison_sheet.cell(1, 1, DISCLAIMER)
    comparison_sheet.cell(1, 1).font = Font(bold=True, color="B42318", size=13)
    comparison_sheet.cell(1, 1).fill = PatternFill("solid", fgColor="FEE4E2")
    comparison_sheet.cell(1, 1).alignment = Alignment(horizontal="center")
    comparison_sheet.append(["对比项", *[item["sku"] for item in candidates]])
    comparison_rows = [
        ("产品名称", lambda item: item["model_name"]),
        ("匹配度", lambda item: f"{item['fit_score']:.1f}"),
        ("适配说明", lambda item: item["match_reason"]),
        (
            "工作区间",
            lambda item: (
                f"流量 {item['flow_min_m3h']:g}–{item['flow_max_m3h']:g} m³/h；"
                f"扬程 {item['head_min_m']:g}–{item['head_max_m']:g} m"
            ),
        ),
        ("过流材质", lambda item: item["wetted_material"]),
        ("电机功率", lambda item: f"{item['motor_power_kw']:g} kW"),
        ("密封形式", lambda item: item["seal_type"]),
        ("参考交期", lambda item: f"{item['lead_time_days']:g} 天"),
        ("参考单价", lambda item: f"¥{item['reference_price_cny']:,.0f}"),
        ("价格说明", lambda item: item["price_note"]),
        ("风险提示", lambda item: "；".join(item.get("risks", [])) or "无新增提示"),
    ]
    for label, getter in comparison_rows:
        comparison_sheet.append([label, *[getter(item) for item in candidates]])
    for cell in comparison_sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5EFF")
        cell.alignment = Alignment(horizontal="center")
    _style_table(comparison_sheet, 2, comparison_sheet.max_row, 1, 1 + len(candidates))
    comparison_sheet.column_dimensions["A"].width = 20
    for column in range(2, 2 + len(candidates)):
        comparison_sheet.column_dimensions[get_column_letter(column)].width = 34

    evidence_sheet = workbook.create_sheet("证据与审批")
    evidence_sheet.append(["知识库证据"])
    evidence_sheet["A1"].font = Font(bold=True, size=13)
    evidence_sheet.append(["知识库", "文档", "片段", "相关度"])
    evidence = recommendation.get("evidence", [])
    if evidence:
        for item in evidence:
            evidence_sheet.append(
                [
                    item.get("kb_name", ""),
                    item.get("doc_name", ""),
                    item.get("content", ""),
                    item.get("score", ""),
                ]
            )
    else:
        evidence_sheet.append(["未配置或未检索到知识库证据", "", "", ""])
    approval_start = evidence_sheet.max_row + 2
    evidence_sheet.cell(approval_start, 1, "审批记录")
    evidence_sheet.cell(approval_start, 1).font = Font(bold=True, size=13)
    evidence_sheet.append(["类型", "决定", "操作人", "意见", "时间"])
    for approval in case.get("approvals", []):
        evidence_sheet.append(
            [
                approval["approval_type"],
                approval["decision"],
                approval["actor"],
                approval["comment"],
                approval["created_at"],
            ]
        )
    _style_table(evidence_sheet, 2, evidence_sheet.max_row, 1, 5)
    evidence_sheet.column_dimensions["A"].width = 22
    evidence_sheet.column_dimensions["B"].width = 28
    evidence_sheet.column_dimensions["C"].width = 72
    evidence_sheet.column_dimensions["D"].width = 22
    evidence_sheet.column_dimensions["E"].width = 28

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
    workbook.save(path)
    return path
