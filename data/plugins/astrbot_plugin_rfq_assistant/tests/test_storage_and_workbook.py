"""Persistence, workflow, catalog, and Excel tests."""

from copy import deepcopy

import pytest
from openpyxl import load_workbook

from data.plugins.astrbot_plugin_rfq_assistant.domain import (
    DEFAULT_PRODUCTS,
    apply_manual_fields,
    build_response_text,
    compute_missing,
    fallback_extract,
    merge_extraction,
    recommend_products,
)
from data.plugins.astrbot_plugin_rfq_assistant.storage import (
    InvalidTransitionError,
    RevisionConflictError,
    RfqStorage,
)
from data.plugins.astrbot_plugin_rfq_assistant.workbook import (
    DISCLAIMER,
    build_case_export,
    build_catalog_template,
    read_catalog,
)


def complete_requirements():
    """Create one complete acceptance-case requirement mapping.

    Returns:
        Normalized explicit pump requirements.
    """
    source = "2 台泵，30 m³/h、30 m、380 V，月底交付。40°C 清水、50 Hz"
    return merge_extraction(None, fallback_extract(source), source)


async def create_selectable_case(storage: RfqStorage, case_id: str = "RFQ-TEST-0001"):
    """Create a complete case with a persisted recommendation snapshot.

    Args:
        storage: Initialized test storage.
        case_id: Deterministic test identifier.

    Returns:
        Full persisted case.
    """
    requirements = complete_requirements()
    missing = compute_missing(requirements)
    version_id, products = await storage.active_catalog()
    selection = recommend_products(products, requirements)
    await storage.create_case(
        case_id=case_id,
        source_text="完整询价",
        requirements=requirements,
        missing=missing,
        response_text=build_response_text(
            case_id, requirements, missing, selection["candidates"]
        ),
        actor="dashboard:sales",
    )
    await storage.save_recommendations(
        case_id=case_id,
        revision_no=1,
        catalog_version_id=version_id,
        candidates=selection["candidates"],
        conflicts=selection["conflicts"],
        evidence=[],
        actor="dashboard:sales",
    )
    return await storage.get_case(case_id)


@pytest.mark.asyncio
async def test_initialize_seeds_versioned_synthetic_catalog(tmp_path) -> None:
    storage = RfqStorage(tmp_path / "rfq.db")

    await storage.initialize()
    version_id, products = await storage.active_catalog()

    assert version_id == 1
    assert [item["sku"] for item in products] == sorted(
        item["sku"] for item in DEFAULT_PRODUCTS
    )
    assert all("虚构演示参考价" in item["price_note"] for item in products)


@pytest.mark.asyncio
async def test_catalog_import_rolls_back_on_duplicate_sku(tmp_path) -> None:
    storage = RfqStorage(tmp_path / "rfq.db")
    await storage.initialize()
    original_version, original_products = await storage.active_catalog()
    invalid_products = [deepcopy(DEFAULT_PRODUCTS[0]), deepcopy(DEFAULT_PRODUCTS[0])]

    with pytest.raises(Exception):
        await storage.import_catalog(invalid_products, "重复 SKU")

    active_version, active_products = await storage.active_catalog()
    assert active_version == original_version
    assert active_products == original_products


@pytest.mark.asyncio
async def test_confirmation_approval_and_export_workflow(tmp_path) -> None:
    storage = RfqStorage(tmp_path / "rfq.db")
    await storage.initialize()
    case = await create_selectable_case(storage)

    with pytest.raises(InvalidTransitionError):
        await storage.approve_solutions(
            case["id"], 1, "dashboard:engineer", "approved"
        )

    case = await storage.confirm_requirements(
        case["id"], 1, "dashboard:sales", "参数已核对"
    )
    assert case["status"] == "awaiting_technical_review"
    case = await storage.approve_solutions(
        case["id"], 1, "dashboard:engineer", "approved", "选型可用"
    )
    assert case["status"] == "approved_for_export"

    case = await storage.mark_exported(case["id"], "dashboard:sales")
    assert case["status"] == "approved_for_export"
    assert case["audit"][0]["event_type"] == "reference_export_generated"


@pytest.mark.asyncio
async def test_new_revision_invalidates_current_approvals_and_detects_conflict(
    tmp_path,
) -> None:
    storage = RfqStorage(tmp_path / "rfq.db")
    await storage.initialize()
    case = await create_selectable_case(storage)
    await storage.confirm_requirements(case["id"], 1, "dashboard:sales")
    await storage.approve_solutions(
        case["id"], 1, "dashboard:engineer", "approved"
    )
    requirements = apply_manual_fields(
        case["revision"]["requirements"], {"temperature_c": 45}
    )

    updated = await storage.update_case(
        case_id=case["id"],
        requirements=requirements,
        missing=compute_missing(requirements),
        response_text="重新选型",
        actor="dashboard:sales",
        expected_revision=1,
        source_text_append="介质温度修订为 45°C",
    )

    assert updated["active_revision"] == 2
    assert updated["status"] == "awaiting_sales_confirmation"
    assert updated["approvals"] == []
    with pytest.raises(RevisionConflictError):
        await storage.update_case(
            case_id=case["id"],
            requirements=requirements,
            missing=[],
            response_text="stale",
            actor="dashboard:sales",
            expected_revision=1,
        )


def test_catalog_template_round_trip_and_strict_columns(tmp_path) -> None:
    path = build_catalog_template(tmp_path / "catalog.xlsx")

    products = read_catalog(path)

    assert len(products) == 5
    assert products[0]["sku"] == "CP-32-125E"
    workbook = load_workbook(path)
    workbook["产品目录"].cell(1, 1, "wrong_header")
    workbook.save(path)
    with pytest.raises(ValueError, match="列与模板不一致"):
        read_catalog(path)


@pytest.mark.asyncio
async def test_approved_export_has_three_sheets_and_no_total_formula(tmp_path) -> None:
    storage = RfqStorage(tmp_path / "rfq.db")
    await storage.initialize()
    case = await create_selectable_case(storage)
    case = await storage.confirm_requirements(
        case["id"], 1, "dashboard:sales", "已核对"
    )
    case = await storage.approve_solutions(
        case["id"], 1, "dashboard:engineer", "approved", "同意"
    )
    output = build_case_export(tmp_path / "选型与参考价格表.xlsx", case)

    workbook = load_workbook(output, data_only=False)

    assert workbook.sheetnames == ["需求确认", "方案对比", "证据与审批"]
    assert workbook["需求确认"]["A1"].value == DISCLAIMER
    assert workbook["方案对比"]["A1"].value == DISCLAIMER
    comparison_values = [
        cell.value
        for row in workbook["方案对比"].iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert {"CP-40-160E", "CP-40-160S", "CP-50-200H"}.issubset(
        set(comparison_values)
    )
    assert not any(
        isinstance(value, str) and value.startswith("=")
        for value in comparison_values
    )
    assert not any("总价" in str(value) or "小计" in str(value) for value in comparison_values)
    approval_values = [
        cell.value
        for row in workbook["证据与审批"].iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert "dashboard:sales" in approval_values
    assert "dashboard:engineer" in approval_values
